# -*- coding: utf-8 -*-
"""
生成《支付对账与差错追踪 · 差异定位辅助》报价单。

用法：python tools/build_quote_xlsx.py
输出：docs/04-方案与报价/报价单_支付对账差异定位辅助.xlsx
"""

from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "04-方案与报价", "报价单_支付对账差异定位辅助.xlsx")

NAVY = "1B3A63"
LIGHT = "F2F5F9"
BAND = "F7F9FC"
ACCENT = "FDF3F2"
LINE = "C9D6E6"
BUDGET = 260000

THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 序号, 工作项, 交付内容, 人天, 单价
ITEMS = [
    ("01", "业务梳理与流程对齐", "六节点流程还原、责任与权限矩阵、指标口径草案", 8, 2000),
    ("02", "数据接入与字段级质控", "四类材料加载、版本登记、Q1–Q7 校验、人工修复队列", 12, 2200),
    ("03", "分层匹配引擎", "L1/L2/L3 + 合并入账拆分 + 退款事件反查，含置信度分层", 15, 2400),
    ("04", "差异原因码与证据链", "10 类原因码、证据引用、可查原因完整性校验", 10, 2200),
    ("05", "人工兜底与工单", "四类触发判定、岗位池分派、SLA 与超时升级", 8, 2000),
    ("06", "审计留痕与权限控制", "六要素留痕、哈希链校验、岗位最小化视图", 8, 2200),
    ("07", "复核台与指标看板", "差异下钻、修复队列、工单看板、双口径指标", 12, 2000),
    ("08", "离线回放与验收支撑", "单渠道回放、连续 5 日复测、验收报告", 10, 2200),
    ("09", "培训与文档交付", "操作手册、运维手册、两场岗位培训", 6, 1800),
    ("10", "并行试点支持", "并行期现场支持、误判复盘、规则微调", 10, 1800),
]

PAY_STAGES = [
    ("P0 合同签订", "第 1 周", 0.30, "口径定义草案、权限矩阵"),
    ("P2 匹配与定位交付", "第 6 周", 0.40, "单渠道回放报告、原因码全量样本"),
    ("P4 验收通过", "第 10 周", 0.20, "验收报告、连续 5 日复测结果"),
    ("质保期满", "验收后 60 天", 0.10, "质保期问题清零记录"),
]


def money(v):
    return round(v, 2)


def build():
    wb = Workbook()

    # ==================================================== Sheet1 报价明细
    ws = wb.active
    ws.title = "报价明细"
    ws.sheet_view.showGridLines = False

    widths = [6, 26, 46, 8, 11, 13, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells("A1:G1")
    c = ws["A1"]
    c.value = "支付对账与差错追踪 · 差异定位辅助（只读）  报 价 单"
    c.font = Font(name="微软雅黑", size=15, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 40

    ws.merge_cells("A2:G2")
    c = ws["A2"]
    c.value = ("报价单位：恒川支付运营数字化小组        报价日期：2026-10-29        "
               "报价编号：HC-RECON-2026-001        有效期：30 天")
    c.font = Font(name="微软雅黑", size=9, color="5D6D84")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 22

    headers = ["序号", "工作项", "交付内容", "人天", "单价（元/人天）", "小计（元）", "备注"]
    hr = 4
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=hr, column=i, value=h)
        c.font = Font(name="微软雅黑", size=9.5, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[hr].height = 30

    notes = {
        "01": "P0 阶段，不通过不进入开发",
        "02": "决定匹配率上限的关键环节",
        "03": "不含自动调账等资金动作",
        "04": "差异必须可回指原记录",
        "05": "四类触发缺一不可",
        "06": "哈希链，可验篡改",
        "07": "双口径同时披露",
        "08": "合成样本不作为达标证据",
        "09": "面向 4 个岗位",
        "10": "并行期不替换现有业务",
    }

    r = hr + 1
    subtotal = 0
    days_total = 0
    for idx, (no, name, deliver, days, price) in enumerate(ITEMS):
        amount = days * price
        subtotal += amount
        days_total += days
        vals = [no, name, deliver, days, price, amount, notes[no]]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=ci, value=v)
            c.border = BORDER
            c.font = Font(name="微软雅黑", size=9.5, bold=(ci == 2))
            if ci in (4, 5, 6):
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = "#,##0" if ci != 4 else "0"
            elif ci in (1,):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
            if idx % 2 == 1:
                c.fill = PatternFill("solid", fgColor=BAND)
        ws.row_dimensions[r].height = 30
        r += 1

    # 小计
    def summary_row(row, label, value_formula, note="", bold=True, fill=LIGHT, size=9.5):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        c = ws.cell(row=row, column=1, value=label)
        c.font = Font(name="微软雅黑", size=size, bold=bold)
        c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        c.fill = PatternFill("solid", fgColor=fill)
        c.border = BORDER
        for ci in range(2, 6):
            ws.cell(row=row, column=ci).fill = PatternFill("solid", fgColor=fill)
            ws.cell(row=row, column=ci).border = BORDER
        v = ws.cell(row=row, column=6, value=value_formula)
        v.font = Font(name="微软雅黑", size=size, bold=bold)
        v.alignment = Alignment(horizontal="right", vertical="center")
        v.number_format = "#,##0.00"
        v.fill = PatternFill("solid", fgColor=fill)
        v.border = BORDER
        n = ws.cell(row=row, column=7, value=note)
        n.font = Font(name="微软雅黑", size=8.5, color="5D6D84")
        n.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        n.fill = PatternFill("solid", fgColor=fill)
        n.border = BORDER
        ws.row_dimensions[row].height = 26

    summary_row(r, "实施小计（%d 人天）" % days_total, subtotal, "")
    sub_row = r
    r += 1
    pm = money(subtotal * 0.08)
    summary_row(r, "项目管理与质量保障（实施小计 × 8%）", pm, "进度、评审、缺陷管理")
    pm_row = r
    r += 1
    base = money(subtotal + pm)
    summary_row(r, "不含税合计", base, "")
    base_row = r
    r += 1
    tax = money(base * 0.06)
    summary_row(r, "增值税（6%，专票）", tax, "")
    tax_row = r
    r += 1
    total = money(base + tax)
    summary_row(r, "报价总计（含税）", total, "人民币", bold=True, fill="E8EEF6", size=11)
    total_row = r
    ws.row_dimensions[total_row].height = 32
    r += 1
    summary_row(r, "首期预算（f03）", BUDGET, "对照财务口径", fill="FFFFFF")
    r += 1
    summary_row(r, "预算余量", money(BUDGET - total),
                "占预算 %.1f%%，用于范围变更与复测" % ((BUDGET - total) / BUDGET * 100),
                fill=ACCENT)
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ws.cell(row=r, column=1, value=(
        "说明：① 报价为单渠道（CH-WEIXIN-A）离线回放范围，不含多渠道路由、实时对账、自动调账；"
        "② 人天按 8 小时/人天计；③ 不含硬件、第三方商业软件许可与渠道侧接口改造成本；"
        "④ 合成样本验证结论不构成 99.9% 达标证据，达标以真实账单回放结果为准。"))
    c.font = Font(name="微软雅黑", size=8.5, color="5D6D84")
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[r].height = 46

    # ==================================================== Sheet2 付款与工期
    ws2 = wb.create_sheet("付款与工期")
    ws2.sheet_view.showGridLines = False
    for i, w in enumerate([22, 18, 12, 16, 44], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws2.merge_cells("A1:E1")
    c = ws2["A1"]
    c.value = "付款节点与工期安排"
    c.font = Font(name="微软雅黑", size=13, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws2.row_dimensions[1].height = 34

    heads = ["付款节点", "时间", "比例", "金额（元）", "对应交付物"]
    for i, h in enumerate(heads, 1):
        c = ws2.cell(row=3, column=i, value=h)
        c.font = Font(name="微软雅黑", size=9.5, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
    ws2.row_dimensions[3].height = 26

    paid = 0.0
    for i, (stage, when, ratio, deliver) in enumerate(PAY_STAGES):
        if i == len(PAY_STAGES) - 1:
            amt = money(total - paid)          # 末笔用差额，保证合计精确
        else:
            amt = money(total * ratio)
            paid += amt
        vals = [stage, when, "%.0f%%" % (ratio * 100), amt, deliver]
        for ci, v in enumerate(vals, 1):
            c = ws2.cell(row=4 + i, column=ci, value=v)
            c.border = BORDER
            c.font = Font(name="微软雅黑", size=9.5, bold=(ci == 1))
            if ci == 4:
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif ci in (2, 3):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
            if i % 2 == 1:
                c.fill = PatternFill("solid", fgColor=BAND)
        ws2.row_dimensions[4 + i].height = 30

    ws2.merge_cells("A9:E9")
    c = ws2["A9"]
    c.value = ("合计：%.2f 元        总工期：10 周（含 P0 口径对齐 1 周）        "
               "质保期：验收后 60 天" % total)
    c.font = Font(name="微软雅黑", size=10, bold=True)
    c.fill = PatternFill("solid", fgColor=LIGHT)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c.border = BORDER
    ws2.row_dimensions[9].height = 30

    ws2.merge_cells("A11:E11")
    c = ws2["A11"]
    c.value = "关键前提：P0 口径对齐（99.9% 分母定义、排除规则、可查原因字段清单）未签核前，不进入 P1 开发。"
    c.font = Font(name="微软雅黑", size=9, color="B42318")
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    ws2.row_dimensions[11].height = 26

    # ==================================================== Sheet3 报价说明
    ws3 = wb.create_sheet("报价说明")
    ws3.sheet_view.showGridLines = False
    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["B"].width = 92

    ws3.merge_cells("A1:B1")
    c = ws3["A1"]
    c.value = "报价范围与商务条款"
    c.font = Font(name="微软雅黑", size=13, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws3.row_dimensions[1].height = 34

    rows = [
        ("报价包含", "① 单渠道离线回放所需的全部开发、配置、测试；② 四类材料（渠道结算文件、内部流水、"
                     "退款事件、手续费规则）的只读接入；③ 10 类差异原因码与证据链；④ 四类转人工兜底与工单；"
                     "⑤ 哈希链审计；⑥ 复核台与双口径指标；⑦ 两场岗位培训与全套文档。"),
        ("报价不含", "① 服务器、存储、网络等硬件；② 第三方商业软件许可（如商业对账产品、OCR、图数据库）；"
                     "③ 渠道侧或财务总账侧的接口改造（按实结算）；④ 多渠道路由与实时对账；"
                     "⑤ 自动调账、自动平账等任何资金动作能力；⑥ 服务侧解释工单系统。"),
        ("计价方式", "按人天计价，人天单价 1,800–2,400 元，按岗位与工作项复杂度分档；"
                     "项目管理与质量保障按实施小计 8% 计取。"),
        ("税费", "增值税税率 6%，开具增值税专用发票。"),
        ("工期", "10 周。其中 P0 口径对齐 1 周为前置条件，不通过则后续工期顺延。"),
        ("验收依据", "以《解决方案设计》第十二节验收清单为准；准确率以真实账单回放结果计，"
                     "合成样本结果不作为达标证据。"),
        ("变更", "范围变更按人天单价另计，单次变更超过 5 人天需重新评审工期。"),
        ("付款", "30% / 40% / 20% / 10%（质保期满），见「付款与工期」页。"),
        ("质保", "验收后 60 天内免费修复缺陷；新增需求不属于质保范围。"),
        ("有效期", "自报价日起 30 天。"),
    ]
    for i, (k, v) in enumerate(rows):
        r = 3 + i
        c1 = ws3.cell(row=r, column=1, value=k)
        c1.font = Font(name="微软雅黑", size=9.5, bold=True)
        c1.fill = PatternFill("solid", fgColor=LIGHT)
        c1.alignment = Alignment(horizontal="left", vertical="top", indent=1)
        c1.border = BORDER
        c2 = ws3.cell(row=r, column=2, value=v)
        c2.font = Font(name="微软雅黑", size=9.5)
        c2.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
        c2.border = BORDER
        ws3.row_dimensions[r].height = 52 if len(v) > 60 else 26

    # 打印设置
    for sh in (ws, ws2, ws3):
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 1 if sh is ws else 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.page_margins.left = sh.page_margins.right = 0.4
        sh.page_margins.top = sh.page_margins.bottom = 0.5

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("已生成：%s" % OUT)
    print("实施小计 %.2f ｜ 含税合计 %.2f ｜ 预算余量 %.2f" % (subtotal, total, BUDGET - total))


if __name__ == "__main__":
    build()
